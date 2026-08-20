#!/usr/bin/env python3
"""Mantiene el brochure de empresas coherente con la página /empresas.

Qué hace sobre public/brochure-delistars-empresas.docx:
  1. Quita "Lo que dicen nuestros clientes" — esos testimonios se eliminaron de
     la página (no había comentarios relevantes) y el brochure se quedó atrás.
  2. Inserta la sección de ingredientes leyendo src/data/ingredientes.json, el
     mismo archivo que alimenta la página y la ficha descargable.
  3. Exporta el .docx a PDF con Word, que es de donde sale el brochure que se
     descarga desde la página.

    python scripts/actualizar-brochure.py

El script es idempotente: si la sección de ingredientes ya está, la reemplaza.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

RAIZ  = Path(__file__).resolve().parent.parent
DATOS = RAIZ / "src" / "data" / "ingredientes.json"
DOCX  = RAIZ / "public" / "brochure-delistars-empresas.docx"
PDF   = RAIZ / "public" / "brochure-delistars-empresas.pdf"

# Paleta tomada del propio documento, para que lo insertado no desentone.
CHERRY = RGBColor(0xEA, 0x33, 0x29)   # títulos de sección
GRIS   = RGBColor(0x4A, 0x4A, 0x4A)   # texto de apoyo
COAL   = RGBColor(0x26, 0x26, 0x26)   # texto principal
MINT   = RGBColor(0x18, 0x93, 0x86)   # etiqueta de viñeta

TITULO_INGREDIENTES = "¿Qué lleva cada producto?"
BLOQUE_TESTIMONIOS  = "Lo que dicen nuestros clientes"
ANCLA_FINAL         = "Cómo funciona"


def parrafos(doc):
    return [Paragraph(el, doc) for el in doc.element.body.iterchildren()
            if el.tag.endswith("}p")]


def buscar(doc, texto):
    for p in parrafos(doc):
        if p.text.strip().lower() == texto.lower():
            return p
    return None


def escribir(p, texto, *, size=11, bold=False, color=GRIS):
    r = p.add_run(texto)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color
    return r


def nuevo_antes(ancla, doc):
    """Crea un párrafo vacío justo antes del ancla y lo devuelve."""
    p = doc.add_paragraph()
    ancla._p.addprevious(p._p)
    return p


def quitar_testimonios(doc):
    """Borra el encabezado y todo lo que cuelga de él hasta el siguiente título."""
    cuerpo = list(doc.element.body.iterchildren())
    inicio = None
    for i, el in enumerate(cuerpo):
        if el.tag.endswith("}p"):
            if Paragraph(el, doc).text.strip().lower() == BLOQUE_TESTIMONIOS.lower():
                inicio = i
                break
    if inicio is None:
        print("· testimonios: ya no estaban")
        return 0

    # Se corta hasta el siguiente título de sección (bold 17) sin incluirlo.
    fin = len(cuerpo)
    for j in range(inicio + 1, len(cuerpo)):
        el = cuerpo[j]
        if not el.tag.endswith("}p"):
            continue
        p = Paragraph(el, doc)
        r = p.runs[0] if p.runs else None
        if p.text.strip() and r is not None and r.bold and r.font.size and r.font.size.pt >= 15:
            fin = j
            break

    for el in cuerpo[inicio:fin]:
        el.getparent().remove(el)
    print(f"· testimonios: {fin - inicio} elementos eliminados")
    return fin - inicio


def quitar_ingredientes_previos(doc):
    """Para que el script se pueda correr varias veces sin duplicar."""
    cuerpo = list(doc.element.body.iterchildren())
    inicio = None
    for i, el in enumerate(cuerpo):
        if el.tag.endswith("}p") and Paragraph(el, doc).text.strip() == TITULO_INGREDIENTES:
            inicio = i
            break
    if inicio is None:
        return
    fin = len(cuerpo)
    for j in range(inicio + 1, len(cuerpo)):
        el = cuerpo[j]
        if el.tag.endswith("}p") and Paragraph(el, doc).text.strip().lower() == ANCLA_FINAL.lower():
            fin = j
            break
    for el in cuerpo[inicio:fin]:
        el.getparent().remove(el)
    print("· sección de ingredientes anterior: reemplazada")


def insertar_ingredientes(doc, datos):
    ancla = buscar(doc, ANCLA_FINAL)
    if ancla is None:
        sys.exit(f"No encontré el título «{ANCLA_FINAL}» para anclar la sección.")

    escribir(nuevo_antes(ancla, doc), TITULO_INGREDIENTES, size=17, bold=True, color=CHERRY)
    nuevo_antes(ancla, doc)
    escribir(nuevo_antes(ancla, doc),
             "Para que nadie se lleve sorpresas: esto es lo que trae cada uno de nuestros productos.")

    if datos.get("ensalada"):
        escribir(nuevo_antes(ancla, doc), datos["ensaladaNota"], bold=True, color=COAL)

    for fam in datos["familias"]:
        nuevo_antes(ancla, doc)
        p = nuevo_antes(ancla, doc)
        escribir(p, f"●  {fam['titulo']}", bold=True, color=MINT)
        if fam.get("base"):
            escribir(p, "  —  Todos llevan: " + " · ".join(fam["base"]), color=COAL)
        if fam.get("nota"):
            escribir(nuevo_antes(ancla, doc), fam["nota"])
        for prod in fam["productos"]:
            q = nuevo_antes(ancla, doc)
            escribir(q, f"{prod['nombre']}: ", bold=True, color=COAL)
            escribir(q, ", ".join(prod["lleva"]) if prod["lleva"] else "solo la preparación base")

    nuevo_antes(ancla, doc)
    p = nuevo_antes(ancla, doc)
    escribir(p, "Salsas de la casa, incluidas: ", bold=True, color=COAL)
    escribir(p, ", ".join(datos["salsas"]) + ".")
    p = nuevo_antes(ancla, doc)
    escribir(p, "Cebolla a elegir: ", bold=True, color=COAL)
    escribir(p, ", ".join(datos["cebollas"]).lower() + ".")
    escribir(nuevo_antes(ancla, doc),
             "¿Tienes una alergia o una restricción alimentaria? Escríbenos antes de pedir y te "
             "confirmamos la preparación exacta del producto que te interesa.")
    nuevo_antes(ancla, doc)
    print("· ingredientes: sección insertada antes de «Cómo funciona»")


def exportar_pdf():
    """Word es quien generó el PDF original, así que se conserva el mismo look.

    Se hace con PowerShell porque el intérprete de este entorno no trae
    pywin32; Word sí está instalado y responde por COM.
    """
    import subprocess
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{DOCX}'); $d.SaveAs([ref]'{PDF}', [ref]17); "
        "$d.Close([ref]$false); $w.Quit()"
    )
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"No se pudo exportar el PDF con Word:\n{r.stderr}")
    print(f"· PDF exportado: {PDF} ({PDF.stat().st_size / 1024:.0f} KB)")


def main():
    datos = json.loads(DATOS.read_text(encoding="utf-8"))
    doc = Document(str(DOCX))
    quitar_testimonios(doc)
    quitar_ingredientes_previos(doc)
    insertar_ingredientes(doc, datos)
    doc.save(str(DOCX))
    print(f"· DOCX guardado: {DOCX}")
    exportar_pdf()


if __name__ == "__main__":
    main()
